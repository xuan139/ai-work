from __future__ import annotations

import asyncio
import json
import os
import re
from functools import lru_cache
from pathlib import Path
from typing import Any

from app.asr_catalog import get_asr_model
from app.asr_runtime import AsrRuntimeError
from app.asr_service import run_asr_with_audit
from app.db import get_nas_asset, replace_document_chunks, update_nas_asset
from app.embedding_runtime import attach_embeddings
from app.llm_runtime import LlmRuntimeError
from app.notifications import manager
from app.translation_service import translate_transcript_with_audit
from app.video_catalog import get_video_model
from app.video_runtime import VideoRuntimeError, analyze_video, build_video_detection_chunks

BASE_DIR = Path(__file__).resolve().parent.parent
AUDIO_SUFFIXES = {".wav", ".mp3", ".m4a", ".webm", ".ogg", ".flac", ".aac"}
VIDEO_SUFFIXES = {".mp4", ".mov", ".mkv", ".avi", ".webm", ".m4v"}
PDF_SUFFIXES = {".pdf"}
DOCX_SUFFIXES = {".docx"}
IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".webp", ".tif", ".tiff", ".bmp"}
PDF_RENDER_SCALE = 2.0


def classify_asset(path: Path, mime_type: str | None = None) -> str:
    suffix = path.suffix.lower()
    mime = (mime_type or "").lower()
    if suffix in AUDIO_SUFFIXES or mime.startswith("audio/"):
        return "audio"
    if suffix in VIDEO_SUFFIXES or mime.startswith("video/"):
        return "video"
    if suffix in PDF_SUFFIXES or mime == "application/pdf":
        return "pdf"
    if suffix in DOCX_SUFFIXES or "wordprocessingml.document" in mime:
        return "docx"
    if suffix in IMAGE_SUFFIXES or mime.startswith("image/"):
        return "image"
    return "file"


def analyzer_for_category(category: str) -> str:
    return {
        "audio": "Whisper",
        "video": "YOLO",
        "pdf": "RAG Builder",
        "docx": "RAG Builder",
        "image": "PaddleOCR Image RAG",
    }.get(category, "NAS Indexer")


async def process_nas_asset(
    asset_id: int,
    audio_api_key: str | None = None,
    video_api_key: str | None = None,
    translation_api_key: str | None = None,
) -> None:
    asset = get_nas_asset(asset_id)
    if not asset:
        return

    try:
        category = asset["category"]
        path = Path(asset["stored_path"])
        if category == "audio":
            updated = await process_audio_asset(asset, path, audio_api_key, translation_api_key)
        elif category == "video":
            updated = await process_video_asset(asset, path, video_api_key)
        elif category in {"pdf", "docx"}:
            chunks = await asyncio.to_thread(build_rag_chunks, path, category, asset_id)
            embedded = await attach_embeddings(chunks, asset["user_id"])
            replace_document_chunks(asset_id, chunks)
            updated = update_nas_asset(
                asset_id,
                status="completed",
                analyzer="RAG Builder",
                summary=(
                    f"已抽取文件文字並建立 RAG chunks：{len(chunks)} 段；"
                    f"{'Qwen3 Embedding 向量已寫入' if embedded else '向量服務暫時不可用，已排入背景補建'}，"
                    "可在此頁使用 LLM 進行文件問答。"
                ),
                chunk_count=len(chunks),
            )
        elif category == "image":
            chunks = await asyncio.to_thread(build_image_rag_chunks, path)
            embedded = await attach_embeddings(chunks, asset["user_id"])
            replace_document_chunks(asset_id, chunks)
            updated = update_nas_asset(
                asset_id,
                status="completed",
                analyzer="PaddleOCR Image RAG",
                summary=(
                    f"已使用 PaddleOCR 辨識圖片文字並建立 image_ocr chunks：{len(chunks)} 段；"
                    f"{'Qwen3 Embedding 向量已寫入' if embedded else '向量服務暫時不可用，已排入背景補建'}，"
                    "可在此頁依原圖來源使用 LLM 進行 RAG 問答。"
                ),
                chunk_count=len(chunks),
            )
        else:
            updated = update_nas_asset(
                asset_id,
                status="completed",
                analyzer="NAS Indexer",
                summary="檔案已保存並完成 NAS 索引。此類型目前只保留原始檔與中繼資料。",
                chunk_count=0,
            )
        await _notify(asset, updated, "nas_asset_processed")
    except Exception as exc:
        updated = update_nas_asset(
            asset_id,
            status="failed",
            analyzer=asset.get("analyzer"),
            summary=None,
            error_message=str(exc),
        )
        await _notify(asset, updated, "nas_asset_failed")


async def process_audio_asset(
    asset: dict[str, Any],
    path: Path,
    api_key: str | None,
    translation_api_key: str | None,
) -> dict[str, Any] | None:
    config = processor_config(asset)
    model = get_asr_model(config.get("asr_model_id"))
    try:
        result = await run_asr_with_audit(
            path=path,
            model_id=model["id"],
            api_key=api_key,
            user_id=asset["user_id"],
        )
    except AsrRuntimeError as exc:
        local_model = model["id"].startswith("local:")
        return update_nas_asset(
            asset["id"],
            status="needs_model" if local_model else "failed",
            analyzer=model["name"],
            summary=f"已選擇 {model['name']}。{exc}",
            error_message=None if local_model else str(exc),
            chunk_count=0,
        )

    chunks = build_transcript_chunks(result["text"], path, model, result)
    translation_summary = ""
    if config.get("translation_enabled"):
        try:
            translated = await translate_transcript_with_audit(
                text=result["text"],
                target=config.get("translation_target") or "zh-Hant",
                model_id=config.get("translation_model_id") or "local:qwen3-4b",
                api_key=translation_api_key,
                user_id=asset["user_id"],
            )
            chunks.extend(
                build_translation_chunks(
                    translated["text"],
                    path,
                    translated["model"],
                    config.get("translation_target") or "zh-Hant",
                    start_index=len(chunks),
                )
            )
            translation_summary = f"已由 {translated['model']['name']} 完成翻譯；"
        except LlmRuntimeError as exc:
            translation_summary = f"逐字稿已保存，但翻譯失敗：{exc}；"
    embedded = await attach_embeddings(chunks, asset["user_id"])
    replace_document_chunks(asset["id"], chunks)
    return update_nas_asset(
        asset["id"],
        status="completed",
        analyzer=model["name"],
        summary=(
            f"已使用 {model['name']} 完成語音轉文字，建立逐字稿 RAG chunks：{len(chunks)} 段；"
            f"{translation_summary}"
            f"{'Qwen3 Embedding 向量已寫入' if embedded else '向量服務暫時不可用，已排入背景補建'}。"
        ),
        chunk_count=len(chunks),
    )


async def process_video_asset(asset: dict[str, Any], path: Path, api_key: str | None) -> dict[str, Any] | None:
    config = processor_config(asset)
    model = get_video_model(config.get("video_model_id"))
    try:
        result = await asyncio.to_thread(analyze_video, path, model["id"], api_key)
    except VideoRuntimeError as exc:
        return update_nas_asset(
            asset["id"],
            status="needs_model",
            analyzer=model["name"],
            summary=f"已選擇 {model['name']}。{exc}",
            chunk_count=0,
        )

    chunks = build_video_detection_chunks(path, result)
    embedded = await attach_embeddings(chunks, asset["user_id"])
    replace_document_chunks(asset["id"], chunks)
    return update_nas_asset(
        asset["id"],
        status="completed",
        analyzer=model["name"],
        summary=(
            f"{result['summary']} 已建立 video_detection chunks：{len(chunks)} 段；"
            f"{'Qwen3 Embedding 向量已寫入' if embedded else '向量服務暫時不可用，已排入背景補建'}，"
            "可在此頁使用 LLM 查詢影片內容。"
        ),
        chunk_count=len(chunks),
    )


def processor_config(asset: dict[str, Any]) -> dict[str, Any]:
    raw = asset.get("processor_config_json")
    if not raw:
        return {}
    try:
        value = json.loads(raw)
    except json.JSONDecodeError:
        return {}
    return value if isinstance(value, dict) else {}


def build_transcript_chunks(text: str, path: Path, model: dict[str, Any], result: dict[str, Any]) -> list[dict[str, Any]]:
    chunks = chunk_text(text, max_chars=1000, overlap=120)
    return [
        {
            "chunk_index": index,
            "content": chunk,
            "token_estimate": max(1, len(chunk) // 4),
            "page_number": None,
            "chunk_type": "audio_transcript",
            "image_path": None,
            "metadata_json": json.dumps(
                {
                    "source": path.name,
                    "category": "audio",
                    "asr_model_id": model["id"],
                    "asr_model": model["name"],
                    "asr_engine": result["engine"],
                },
                ensure_ascii=False,
            ),
        }
        for index, chunk in enumerate(chunks)
    ]


def build_translation_chunks(
    text: str,
    path: Path,
    model: dict[str, Any],
    target: str,
    *,
    start_index: int,
) -> list[dict[str, Any]]:
    chunks = chunk_text(text, max_chars=1000, overlap=120)
    return [
        {
            "chunk_index": start_index + index,
            "content": chunk,
            "token_estimate": max(1, len(chunk) // 4),
            "page_number": None,
            "chunk_type": "audio_translation",
            "image_path": None,
            "metadata_json": json.dumps(
                {
                    "source": path.name,
                    "category": "audio",
                    "translation_target": target,
                    "translation_model_id": model["id"],
                    "translation_model": model["name"],
                },
                ensure_ascii=False,
            ),
        }
        for index, chunk in enumerate(chunks)
    ]


def build_rag_chunks(path: Path, category: str, asset_id: int | None = None) -> list[dict[str, Any]]:
    if category == "pdf":
        return build_pdf_rag_chunks(path, asset_id)

    text = extract_text(path, category)
    if not text.strip():
        raise RuntimeError("文件沒有可抽取的文字內容，無法建立 RAG chunks")

    chunks = chunk_text(text)
    return [
        {
            "chunk_index": index,
            "content": chunk,
            "token_estimate": max(1, len(chunk) // 4),
            "page_number": None,
            "chunk_type": "text",
            "image_path": None,
            "metadata_json": json.dumps({"source": path.name, "category": category}, ensure_ascii=False),
        }
        for index, chunk in enumerate(chunks)
    ]


def build_image_rag_chunks(path: Path) -> list[dict[str, Any]]:
    try:
        engine = paddle_ocr_engine()
    except ImportError as exc:
        raise RuntimeError("需要安裝 OCR 引擎 PaddleOCR，才能辨識圖片並建立 RAG 索引") from exc

    text = run_paddle_ocr(engine, path).strip()
    if not text:
        raise RuntimeError("圖片未辨識出可建立 RAG 索引的文字內容")

    chunks: list[dict[str, Any]] = []
    for text_chunk in chunk_text(text):
        chunks.append(
            {
                "chunk_index": len(chunks),
                "content": text_chunk,
                "token_estimate": max(1, len(text_chunk) // 4),
                "page_number": None,
                "chunk_type": "image_ocr",
                "image_path": str(path),
                "metadata_json": json.dumps(
                    {
                        "source": path.name,
                        "category": "image",
                        "chunk_type": "image_ocr",
                        "extraction_mode": "paddleocr",
                    },
                    ensure_ascii=False,
                ),
            }
        )
    return chunks


def extract_text(path: Path, category: str) -> str:
    if category == "pdf":
        return extract_pdf_text(path)
    if category == "docx":
        return extract_docx_text(path)
    return ""


def extract_pdf_text(path: Path) -> str:
    pages = extract_pdf_text_pages(path)
    return "\n\n".join(
        f"[Page {page['page_number']}]\n{page['text'].strip()}"
        for page in pages
        if page["text"].strip()
    )


def extract_pdf_text_pages(path: Path) -> list[dict[str, Any]]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("缺少 pypdf 套件，無法抽取 PDF 文字") from exc

    reader = PdfReader(str(path))
    pages = []
    for index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        pages.append({"page_number": index, "text": text.strip()})
    return pages


def build_pdf_rag_chunks(path: Path, asset_id: int | None) -> list[dict[str, Any]]:
    text_pages = extract_pdf_text_pages(path)
    page_images = render_pdf_pages(path, asset_id)
    has_text_layer = any(page["text"].strip() for page in text_pages)

    if has_text_layer:
        page_payloads = [
            {
                "page_number": page["page_number"],
                "text": page["text"],
                "chunk_type": "text_layer",
                "extraction_mode": "pypdf",
            }
            for page in text_pages
            if page["text"].strip()
        ]
    else:
        page_payloads = ocr_pdf_pages(page_images)

    chunks: list[dict[str, Any]] = []
    for page in page_payloads:
        page_number = page["page_number"]
        image_path = page_images.get(page_number)
        for text_chunk in chunk_text(page["text"]):
            chunks.append(
                {
                    "chunk_index": len(chunks),
                    "content": text_chunk,
                    "token_estimate": max(1, len(text_chunk) // 4),
                    "page_number": page_number,
                    "chunk_type": page["chunk_type"],
                    "image_path": str(image_path) if image_path else None,
                    "metadata_json": json.dumps(
                        {
                            "source": path.name,
                            "category": "pdf",
                            "page_number": page_number,
                            "chunk_type": page["chunk_type"],
                            "extraction_mode": page["extraction_mode"],
                        },
                        ensure_ascii=False,
                    ),
                }
            )

    if not chunks:
        raise RuntimeError("PDF 沒有可建立 RAG 的文字內容；若是掃描或圖片型 PDF，請確認 OCR 引擎已安裝並可正常讀取")
    return chunks


def render_pdf_pages(path: Path, asset_id: int | None) -> dict[int, Path]:
    try:
        import pymupdf as fitz
    except ImportError as exc:
        try:
            import fitz
        except ImportError:
            raise RuntimeError("缺少 PyMuPDF 套件，無法渲染 PDF 頁面圖片") from exc

    folder_name = f"{path.stem}_pages" if asset_id is None else f"asset_{asset_id}_pages"
    output_dir = path.parent / folder_name
    output_dir.mkdir(parents=True, exist_ok=True)

    rendered: dict[int, Path] = {}
    with fitz.open(str(path)) as document:
        matrix = fitz.Matrix(PDF_RENDER_SCALE, PDF_RENDER_SCALE)
        for page_index in range(document.page_count):
            page_number = page_index + 1
            output_path = output_dir / f"page-{page_number}.png"
            pixmap = document[page_index].get_pixmap(matrix=matrix, alpha=False)
            pixmap.save(str(output_path))
            rendered[page_number] = output_path
    return rendered


def ocr_pdf_pages(page_images: dict[int, Path]) -> list[dict[str, Any]]:
    try:
        engine = paddle_ocr_engine()
    except ImportError as exc:
        raise RuntimeError("需要安裝 OCR 引擎 PaddleOCR，才能處理掃描或圖片型 PDF") from exc

    pages: list[dict[str, Any]] = []
    for page_number, image_path in page_images.items():
        text = run_paddle_ocr(engine, image_path)
        if text.strip():
            pages.append(
                {
                    "page_number": page_number,
                    "text": text.strip(),
                    "chunk_type": "image_ocr",
                    "extraction_mode": "paddleocr",
                }
            )
    return pages


@lru_cache(maxsize=1)
def paddle_ocr_engine() -> Any:
    configure_ocr_cache()
    from paddleocr import PaddleOCR

    try:
        return PaddleOCR(use_angle_cls=True, lang="ch", enable_mkldnn=False)
    except TypeError:
        try:
            return PaddleOCR(lang="ch", enable_mkldnn=False)
        except TypeError:
            return PaddleOCR(lang="ch")


def configure_ocr_cache() -> None:
    cache_root = BASE_DIR / "storage" / "ocr_cache"
    cache_root.mkdir(parents=True, exist_ok=True)
    os.environ.setdefault("PADDLE_PDX_CACHE_HOME", str(cache_root / "paddlex"))
    os.environ.setdefault("MODELSCOPE_CACHE", str(cache_root / "modelscope"))
    os.environ.setdefault("HF_HOME", str(cache_root / "huggingface"))


def run_paddle_ocr(engine: Any, image_path: Path) -> str:
    if hasattr(engine, "ocr"):
        try:
            result = engine.ocr(str(image_path), cls=True)
        except TypeError:
            result = engine.ocr(str(image_path))
    elif hasattr(engine, "predict"):
        result = engine.predict(str(image_path))
    else:
        raise RuntimeError("PaddleOCR 版本不支援已知的 OCR 呼叫介面")
    return "\n".join(collect_ocr_text(result))


def collect_ocr_text(value: Any) -> list[str]:
    texts: list[str] = []
    if value is None:
        return texts
    if isinstance(value, str):
        return [value.strip()] if value.strip() else []
    if isinstance(value, dict):
        for key in ("rec_texts", "texts"):
            items = value.get(key)
            if isinstance(items, list):
                texts.extend(str(item).strip() for item in items if str(item).strip())
        text = value.get("text")
        if isinstance(text, str) and text.strip():
            texts.append(text.strip())
        for item in value.values():
            if isinstance(item, (list, tuple, dict)):
                texts.extend(collect_ocr_text(item))
        return dedupe_texts(texts)
    if isinstance(value, (list, tuple)):
        if len(value) >= 2 and isinstance(value[1], (list, tuple)) and value[1]:
            first = value[1][0]
            if isinstance(first, str) and first.strip():
                texts.append(first.strip())
        for item in value:
            texts.extend(collect_ocr_text(item))
        return dedupe_texts(texts)
    return texts


def dedupe_texts(texts: list[str]) -> list[str]:
    seen = set()
    unique = []
    for text in texts:
        if text and text not in seen:
            seen.add(text)
            unique.append(text)
    return unique


def extract_docx_text(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("缺少 python-docx 套件，無法抽取 DOCX 文字") from exc

    document = Document(str(path))
    parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def chunk_text(text: str, max_chars: int = 1200, overlap: int = 180) -> list[str]:
    normalized = re.sub(r"\n{3,}", "\n\n", text).strip()
    paragraphs = [part.strip() for part in re.split(r"\n\s*\n", normalized) if part.strip()]
    chunks: list[str] = []
    current = ""

    for paragraph in paragraphs:
        if len(paragraph) > max_chars:
            if current:
                chunks.append(current)
                current = ""
            chunks.extend(split_long_text(paragraph, max_chars, overlap))
            continue

        candidate = f"{current}\n\n{paragraph}".strip() if current else paragraph
        if len(candidate) <= max_chars:
            current = candidate
        else:
            chunks.append(current)
            tail = current[-overlap:] if overlap and len(current) > overlap else ""
            current = f"{tail}\n\n{paragraph}".strip() if tail else paragraph

    if current:
        chunks.append(current)
    return chunks


def split_long_text(text: str, max_chars: int, overlap: int) -> list[str]:
    chunks = []
    start = 0
    while start < len(text):
        end = min(len(text), start + max_chars)
        chunks.append(text[start:end].strip())
        if end == len(text):
            break
        start = max(0, end - overlap)
    return [chunk for chunk in chunks if chunk]


async def _notify(asset: dict[str, Any], updated: dict[str, Any] | None, event_type: str) -> None:
    payload_asset = updated or asset
    await manager.broadcast(
        {
            "type": event_type,
            "asset_id": asset["id"],
            "title": asset["title"],
            "message": payload_asset.get("summary") or payload_asset.get("error_message") or asset["title"],
            "asset": payload_asset,
        }
    )
