from __future__ import annotations

import csv
import hashlib
import json
import math
import re
import uuid
from datetime import date, datetime
from itertools import chain
from pathlib import Path
from typing import Any, Callable, Iterable

import duckdb
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.utils import get_column_letter

from app.db import (
    create_excel_pro_job,
    create_nas_asset,
    excel_pro_job_cancel_requested,
    get_excel_pro_job,
    get_nas_asset,
    get_user_by_id,
    list_excel_pro_jobs,
    list_nas_assets,
    request_excel_pro_job_cancel,
    update_excel_pro_job,
    user_can_read_asset,
)


BASE_DIR = Path(__file__).resolve().parent.parent
NAS_ASSETS_DIR = BASE_DIR / "storage" / "nas_assets"
TABULAR_SUFFIXES = {".xlsx", ".xlsm", ".xltx", ".xltm", ".csv", ".tsv"}
MAX_SOURCES = 12
MAX_TABLE_ROWS = 500_000
MAX_QUERY_ROWS = 500
MAX_EXPORT_ROWS = 200_000
MAX_COLUMNS = 200
SQL_BLOCKED = re.compile(
    r"\b(attach|copy|create|delete|detach|drop|export|import|insert|install|load|pragma|replace|set|update|vacuum|call)\b"
    r"|\b(read_[a-z0-9_]+|[a-z0-9_]+_scan|glob|getenv)\s*\(",
    re.IGNORECASE,
)


class ExcelProError(RuntimeError):
    pass


EXCEL_PRO_TOOLS = [
    {
        "name": "excel_pro_list_sources",
        "title": "List Accessible NAS Spreadsheets",
        "description": "List only spreadsheet assets the current NAS user may read, including uploader, department/group, visibility, and asset ID.",
        "inputSchema": {
            "type": "object",
            "properties": {
                "query": {"type": "string"},
                "limit": {"type": "integer", "minimum": 1, "maximum": 100, "default": 50},
            },
            "additionalProperties": False,
        },
        "annotations": {"readOnlyHint": True, "destructiveHint": False},
    },
    {
        "name": "excel_pro_describe_sources",
        "title": "Describe Spreadsheet Tables",
        "description": "Inspect accessible Excel sheets or CSV/TSV files and return DuckDB table aliases, columns, and citable ranges.",
        "inputSchema": {
            "type": "object",
            "properties": {
                "sources": {"type": "array", "items": {"$ref": "#/$defs/source"}, "minItems": 1, "maxItems": MAX_SOURCES}
            },
            "required": ["sources"],
            "$defs": {"source": {"type": "object", "properties": {"asset_id": {"type": "integer"}, "sheet_names": {"type": "array", "items": {"type": "string"}}}, "required": ["asset_id"], "additionalProperties": False}},
            "additionalProperties": False,
        },
        "annotations": {"readOnlyHint": True, "destructiveHint": False},
    },
    {
        "name": "excel_pro_query",
        "title": "Query Excel and CSV with DuckDB",
        "description": "Run one read-only DuckDB SELECT/WITH query across permitted XLSX, CSV, and TSV sources. Returns rows plus worksheet/range citations.",
        "inputSchema": {
            "type": "object",
            "properties": {
                "sources": {"type": "array", "items": {"type": "object", "properties": {"asset_id": {"type": "integer"}, "sheet_names": {"type": "array", "items": {"type": "string"}}}, "required": ["asset_id"], "additionalProperties": False}, "minItems": 1, "maxItems": MAX_SOURCES},
                "sql": {"type": "string"},
                "max_rows": {"type": "integer", "minimum": 1, "maximum": MAX_QUERY_ROWS, "default": 100},
            },
            "required": ["sources", "sql"],
            "additionalProperties": False,
        },
        "annotations": {"readOnlyHint": True, "destructiveHint": False},
    },
    {
        "name": "excel_pro_start_export",
        "title": "Create XLSX Export Job",
        "description": "Start a background job that writes query results to a new private NAS XLSX, optionally adding a pivot summary and chart. Source files are never modified.",
        "inputSchema": {
            "type": "object",
            "properties": {
                "sources": {"type": "array", "items": {"type": "object", "properties": {"asset_id": {"type": "integer"}, "sheet_names": {"type": "array", "items": {"type": "string"}}}, "required": ["asset_id"]}, "minItems": 1, "maxItems": MAX_SOURCES},
                "sql": {"type": "string"}, "output_name": {"type": "string"},
                "pivot": {"type": "object"}, "chart": {"type": "object"},
            },
            "required": ["sources", "sql", "output_name"],
            "additionalProperties": False,
        },
        "annotations": {"readOnlyHint": False, "destructiveHint": False},
    },
    {
        "name": "excel_pro_job_status",
        "title": "Get Excel Export Progress",
        "description": "Read the current user's background Excel export progress and output download location.",
        "inputSchema": {"type": "object", "properties": {"job_id": {"type": "integer"}}, "required": ["job_id"], "additionalProperties": False},
        "annotations": {"readOnlyHint": True, "destructiveHint": False},
    },
]


def list_sources(user: dict[str, Any], query: str = "", limit: int = 50) -> dict[str, Any]:
    assets = list_nas_assets(user_id=int(user["id"]), role=str(user["role"]), q=query or None)
    rows = [_source_summary(asset) for asset in assets if _asset_path(asset).suffix.lower() in TABULAR_SUFFIXES]
    return {"count": min(len(rows), limit), "sources": rows[: max(1, min(limit, 100))]}


def describe_sources(user: dict[str, Any], source_specs: Any) -> dict[str, Any]:
    specs = _normalize_sources(source_specs)
    sources = []
    for spec in specs:
        asset, path = _authorized_asset(user, spec["asset_id"])
        sources.extend(_inspect_asset(asset, path, spec.get("sheet_names")))
    return {"source_count": len(specs), "table_count": len(sources), "tables": sources}


def run_query(user: dict[str, Any], source_specs: Any, sql: str, max_rows: int = 100) -> dict[str, Any]:
    clean_sql = _validate_sql(sql)
    limit = max(1, min(int(max_rows), MAX_QUERY_ROWS))
    connection = duckdb.connect(":memory:")
    try:
        tables = _load_sources(connection, user, _normalize_sources(source_specs))
        _lock_external_access(connection)
        cursor = connection.execute(f"SELECT * FROM ({clean_sql}) AS excel_pro_query LIMIT {limit + 1}")
        columns = [item[0] for item in cursor.description or []]
        raw_rows = cursor.fetchall()
        truncated = len(raw_rows) > limit
        rows = [[_json_value(value) for value in row] for row in raw_rows[:limit]]
        cited = _used_citations(clean_sql, tables)
        return {
            "engine": f"DuckDB {duckdb.__version__}",
            "columns": columns,
            "rows": rows,
            "row_count": len(rows),
            "truncated": truncated,
            "tables": tables,
            "citations": cited,
        }
    finally:
        connection.close()


def enqueue_export(user: dict[str, Any], request: dict[str, Any]) -> dict[str, Any]:
    normalized = _normalized_export_request(user, request)
    job = create_excel_pro_job(user_id=int(user["id"]), request_json=json.dumps(normalized, ensure_ascii=False))
    return serialize_job(job)


def job_status(user: dict[str, Any], job_id: int) -> dict[str, Any]:
    job = get_excel_pro_job(job_id, user_id=int(user["id"]), role=str(user["role"]))
    if not job:
        raise ExcelProError("Excel export job not found")
    return serialize_job(job)


def jobs_for_user(user: dict[str, Any], limit: int = 50) -> list[dict[str, Any]]:
    return [serialize_job(job) for job in list_excel_pro_jobs(user_id=int(user["id"]), role=str(user["role"]), limit=limit)]


def cancel_job(user: dict[str, Any], job_id: int) -> dict[str, Any]:
    job = request_excel_pro_job_cancel(job_id, user_id=int(user["id"]), role=str(user["role"]))
    if not job:
        raise ExcelProError("Excel export job not found")
    return serialize_job(job)


def serialize_job(job: dict[str, Any]) -> dict[str, Any]:
    result = _json_object(job.get("result_json"))
    output_asset_id = job.get("output_asset_id")
    return {
        "id": job["id"], "status": job["status"], "progress": int(job.get("progress") or 0),
        "stage": job.get("stage"), "error_message": job.get("error_message"),
        "created_at": job.get("created_at"), "updated_at": job.get("updated_at"),
        "output_asset_id": output_asset_id, "output_filename": job.get("output_filename"),
        "download_url": f"/api/excel-pro/jobs/{job['id']}/download" if output_asset_id and job["status"] == "completed" else None,
        "result": result,
    }


def execute_export_job(job: dict[str, Any], progress: Callable[[int, str], None]) -> dict[str, Any]:
    request = _json_object(job.get("request_json"))
    user = get_user_by_id(int(job["user_id"]))
    if not user or not user.get("is_active", 1):
        raise ExcelProError("The requesting user is no longer active")
    _check_cancel(job["id"])
    progress(8, "checking_permissions")
    clean_sql = _validate_sql(request["sql"])
    specs = _normalize_sources(request["sources"])
    source_hashes = {spec["asset_id"]: _sha256(_authorized_asset(user, spec["asset_id"])[1]) for spec in specs}
    connection = duckdb.connect(":memory:")
    try:
        progress(20, "loading_sources")
        tables = _load_sources(connection, user, specs, lambda done, total: progress(20 + int(30 * done / max(total, 1)), "loading_sources"))
        _lock_external_access(connection)
        _check_cancel(job["id"])
        progress(55, "running_query")
        connection.execute(f"CREATE TABLE excel_pro_result AS SELECT * FROM ({clean_sql}) AS excel_pro_export LIMIT {MAX_EXPORT_ROWS + 1}")
        count = int(connection.execute("SELECT COUNT(*) FROM excel_pro_result").fetchone()[0])
        if count > MAX_EXPORT_ROWS:
            raise ExcelProError(f"Export exceeds {MAX_EXPORT_ROWS:,} rows; narrow the query")
        cursor = connection.execute("SELECT * FROM excel_pro_result")
        result_columns = [item[0] for item in cursor.description or []]
        result_rows = cursor.fetchall()
        progress(68, "building_workbook")
        workbook = Workbook()
        result_sheet = workbook.active
        result_sheet.title = "Query Result"
        _write_rows(result_sheet, result_columns, result_rows)
        pivot_sheet = _add_pivot_sheet(workbook, connection, result_columns, request.get("pivot"))
        progress(78, "adding_chart")
        _add_chart(workbook, result_sheet, pivot_sheet, request.get("chart"))
        citations = _used_citations(clean_sql, tables)
        _add_sources_sheet(workbook, clean_sql, citations)
        _check_cancel(job["id"])
        progress(88, "saving_output")
        NAS_ASSETS_DIR.mkdir(parents=True, exist_ok=True)
        filename = _output_filename(request["output_name"])
        target = NAS_ASSETS_DIR / f"excel-pro-{uuid.uuid4().hex}.xlsx"
        workbook.save(target)
        workbook.close()
        for asset_id, before in source_hashes.items():
            after = _sha256(_authorized_asset(user, asset_id)[1])
            if before != after:
                target.unlink(missing_ok=True)
                raise ExcelProError("Source file changed during export; output was discarded")
        asset = create_nas_asset(
            user_id=int(job["user_id"]), category="document", title=Path(filename).stem,
            original_filename=filename, stored_path=str(target),
            mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            file_size=target.stat().st_size, status="completed", analyzer="excel-pro",
            processor_config_json=json.dumps({"job_id": job["id"], "source_asset_ids": list(source_hashes)}, ensure_ascii=False),
            source_type="excel_pro_export", visibility="private", scan_status="clean", scan_engine="generated",
            content_sha256=_sha256(target),
        )
        return {"row_count": count, "citation_count": len(citations), "output_asset_id": asset["id"], "filename": filename}
    finally:
        connection.close()


def handle_excel_pro_mcp_request(payload: dict[str, Any], user: dict[str, Any] | None) -> dict[str, Any] | None:
    method = str(payload.get("method") or "")
    request_id = payload.get("id")
    if method == "notifications/initialized":
        return None
    if method == "initialize":
        return _rpc_result(request_id, {"protocolVersion": "2025-06-18", "capabilities": {"tools": {"listChanged": False}}, "serverInfo": {"name": "NAS Excel Pro MCP", "version": "1.0.0"}, "instructions": "Permission-aware DuckDB spreadsheet queries and non-destructive XLSX exports."})
    if method == "ping":
        return _rpc_result(request_id, {})
    if method == "tools/list":
        return _rpc_result(request_id, {"tools": EXCEL_PRO_TOOLS})
    if method != "tools/call":
        return _rpc_error(request_id, -32601, f"Method not found: {method}")
    if not user:
        return _rpc_error(request_id, -32001, "A signed AI Work user delegation is required")
    params = payload.get("params") if isinstance(payload.get("params"), dict) else {}
    name = str(params.get("name") or "")
    args = params.get("arguments") if isinstance(params.get("arguments"), dict) else {}
    try:
        if name == "excel_pro_list_sources":
            output = list_sources(user, str(args.get("query") or ""), int(args.get("limit") or 50))
        elif name == "excel_pro_describe_sources":
            output = describe_sources(user, args.get("sources"))
        elif name == "excel_pro_query":
            output = run_query(user, args.get("sources"), str(args.get("sql") or ""), int(args.get("max_rows") or 100))
        elif name == "excel_pro_start_export":
            output = enqueue_export(user, args)
        elif name == "excel_pro_job_status":
            output = job_status(user, int(args.get("job_id") or 0))
        else:
            raise ExcelProError(f"Unknown Excel Pro tool: {name}")
    except (ValueError, ExcelProError) as exc:
        return _rpc_error(request_id, -32602, str(exc))
    return _rpc_result(request_id, {"content": [{"type": "text", "text": json.dumps(output, ensure_ascii=False)}], "structuredContent": output, "isError": False})


def _authorized_asset(user: dict[str, Any], asset_id: int) -> tuple[dict[str, Any], Path]:
    asset = get_nas_asset(int(asset_id))
    if not asset or not user_can_read_asset(int(asset_id), user_id=int(user["id"]), role=str(user["role"])):
        raise ExcelProError(f"Spreadsheet asset {asset_id} was not found or is not permitted")
    path = _asset_path(asset)
    if int(asset.get("is_current", 1)) != 1 or path.suffix.lower() not in TABULAR_SUFFIXES or not path.is_file():
        raise ExcelProError(f"Asset {asset_id} is not an available XLSX/CSV/TSV file")
    return asset, path


def _asset_path(asset: dict[str, Any]) -> Path:
    return Path(str(asset.get("stored_path") or ""))


def _source_summary(asset: dict[str, Any]) -> dict[str, Any]:
    return {"asset_id": asset["id"], "title": asset["title"], "filename": asset["original_filename"], "uploader": asset.get("owner_username"), "uploaded_at": asset.get("created_at"), "visibility": asset.get("visibility"), "department_group_id": asset.get("owner_group_id"), "file_size": asset.get("file_size")}


def _normalize_sources(value: Any) -> list[dict[str, Any]]:
    if not isinstance(value, list) or not value or len(value) > MAX_SOURCES:
        raise ExcelProError(f"sources must contain 1 to {MAX_SOURCES} entries")
    result = []
    seen = set()
    for item in value:
        if not isinstance(item, dict) or not item.get("asset_id"):
            raise ExcelProError("Each source requires asset_id")
        asset_id = int(item["asset_id"])
        if asset_id in seen:
            continue
        sheets = item.get("sheet_names")
        if sheets is not None and (not isinstance(sheets, list) or not all(isinstance(name, str) for name in sheets)):
            raise ExcelProError("sheet_names must be an array of names")
        result.append({"asset_id": asset_id, "sheet_names": sheets or None})
        seen.add(asset_id)
    return result


def _inspect_asset(asset: dict[str, Any], path: Path, requested: list[str] | None) -> list[dict[str, Any]]:
    if path.suffix.lower() in {".csv", ".tsv"}:
        delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
        with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as stream:
            reader = csv.reader(stream, delimiter=delimiter)
            header = next(reader, [])[:MAX_COLUMNS]
            count = sum(1 for _ in reader)
        return [_table_meta(asset, "data", header, count, 1)]
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        output = []
        for sheet in workbook.worksheets:
            if requested and sheet.title not in requested:
                continue
            header = list(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ()))[:MAX_COLUMNS]
            output.append(_table_meta(asset, sheet.title, header, max(0, sheet.max_row - 1), 1))
        if requested and len(output) != len(set(requested)):
            raise ExcelProError(f"One or more worksheets do not exist in {asset['original_filename']}")
        return output
    finally:
        workbook.close()


def _table_meta(asset: dict[str, Any], sheet: str, header: Iterable[Any], row_count: int, first_data_row: int) -> dict[str, Any]:
    columns = _column_names(list(header))
    table = _table_name(int(asset["id"]), sheet)
    last_column = get_column_letter(max(1, len(columns)))
    last_row = max(first_data_row + 1, row_count + first_data_row)
    return {**_source_summary(asset), "sheet_name": sheet, "table_name": table, "columns": columns, "row_count": row_count, "range": f"{sheet}!A{first_data_row + 1}:{last_column}{last_row}"}


def _load_sources(connection: duckdb.DuckDBPyConnection, user: dict[str, Any], specs: list[dict[str, Any]], callback: Callable[[int, int], None] | None = None) -> list[dict[str, Any]]:
    tables = []
    for index, spec in enumerate(specs, start=1):
        asset, path = _authorized_asset(user, spec["asset_id"])
        if path.suffix.lower() in {".csv", ".tsv"}:
            tables.extend(_load_delimited(connection, asset, path))
        else:
            tables.extend(_load_workbook(connection, asset, path, spec.get("sheet_names")))
        if callback:
            callback(index, len(specs))
    return tables


def _load_delimited(connection: duckdb.DuckDBPyConnection, asset: dict[str, Any], path: Path) -> list[dict[str, Any]]:
    table = _table_name(int(asset["id"]), "data")
    delim = "\t" if path.suffix.lower() == ".tsv" else ","
    connection.execute(f'CREATE TABLE "{_quote(table)}" AS SELECT * FROM read_csv(?, header=true, delim=?, sample_size=50000)', [str(path), delim])
    count = int(connection.execute(f'SELECT COUNT(*) FROM "{_quote(table)}"').fetchone()[0])
    if count > MAX_TABLE_ROWS:
        raise ExcelProError(f"{asset['original_filename']} exceeds the {MAX_TABLE_ROWS:,}-row demo limit")
    columns = [row[0] for row in connection.execute(f'DESCRIBE "{_quote(table)}"').fetchall()][:MAX_COLUMNS]
    meta = _table_meta(asset, "data", columns, count, 1)
    meta["columns"] = columns
    return [meta]


def _load_workbook(connection: duckdb.DuckDBPyConnection, asset: dict[str, Any], path: Path, requested: list[str] | None) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    tables = []
    try:
        for sheet in workbook.worksheets:
            if requested and sheet.title not in requested:
                continue
            iterator = sheet.iter_rows(values_only=True)
            header = list(next(iterator, ()))[:MAX_COLUMNS]
            columns = _column_names(header)
            if not columns:
                continue
            sample = []
            for _ in range(200):
                try:
                    sample.append(list(next(iterator))[: len(columns)])
                except StopIteration:
                    break
            types = [_duck_type([row[i] if i < len(row) else None for row in sample]) for i in range(len(columns))]
            table = _table_name(int(asset["id"]), sheet.title)
            definition = ", ".join(f'"{_quote(name)}" {kind}' for name, kind in zip(columns, types))
            connection.execute(f'CREATE TABLE "{_quote(table)}" ({definition})')
            insert = f'INSERT INTO "{_quote(table)}" VALUES ({", ".join("?" for _ in columns)})'
            total = 0
            batch = []
            for row in chain(sample, iterator):
                if total >= MAX_TABLE_ROWS:
                    raise ExcelProError(f"{asset['original_filename']} / {sheet.title} exceeds the {MAX_TABLE_ROWS:,}-row demo limit")
                batch.append([_coerce(row[i] if i < len(row) else None, types[i]) for i in range(len(columns))])
                total += 1
                if len(batch) >= 1000:
                    connection.executemany(insert, batch)
                    batch.clear()
            if batch:
                connection.executemany(insert, batch)
            meta = _table_meta(asset, sheet.title, columns, total, 1)
            meta["columns"] = columns
            tables.append(meta)
        if requested and {item["sheet_name"] for item in tables} != set(requested):
            raise ExcelProError(f"One or more worksheets do not exist in {asset['original_filename']}")
        return tables
    finally:
        workbook.close()


def _duck_type(values: list[Any]) -> str:
    actual = [value for value in values if value is not None and value != ""]
    if not actual:
        return "VARCHAR"
    if all(isinstance(value, bool) for value in actual):
        return "BOOLEAN"
    if all(isinstance(value, int) and not isinstance(value, bool) for value in actual):
        return "BIGINT"
    if all(isinstance(value, (int, float)) and not isinstance(value, bool) for value in actual):
        return "DOUBLE"
    if all(isinstance(value, (date, datetime)) for value in actual):
        return "TIMESTAMP"
    return "VARCHAR"


def _coerce(value: Any, kind: str) -> Any:
    if value is None or value == "":
        return None
    if kind == "VARCHAR":
        return str(value)
    if kind == "TIMESTAMP" and isinstance(value, date) and not isinstance(value, datetime):
        return datetime.combine(value, datetime.min.time())
    try:
        if kind == "BIGINT": return int(value)
        if kind == "DOUBLE": return float(value)
        if kind == "BOOLEAN": return bool(value)
    except (TypeError, ValueError):
        return None
    return value


def _validate_sql(sql: str) -> str:
    clean = str(sql or "").strip().rstrip(";").strip()
    if not clean or not re.match(r"^(select|with)\b", clean, re.IGNORECASE):
        raise ExcelProError("Only one read-only SELECT or WITH query is allowed")
    if ";" in clean or SQL_BLOCKED.search(clean) or "--" in clean or "/*" in clean:
        raise ExcelProError("SQL contains a blocked statement, function, or comment")
    return clean


def _lock_external_access(connection: duckdb.DuckDBPyConnection) -> None:
    connection.execute("SET enable_external_access = false")


def _used_citations(sql: str, tables: list[dict[str, Any]]) -> list[dict[str, Any]]:
    folded = sql.casefold()
    matched = [table for table in tables if table["table_name"].casefold() in folded]
    selected = matched or tables
    return [{key: table.get(key) for key in ("asset_id", "title", "filename", "uploader", "sheet_name", "range", "table_name", "visibility", "department_group_id")} for table in selected]


def _normalized_export_request(user: dict[str, Any], request: dict[str, Any]) -> dict[str, Any]:
    specs = _normalize_sources(request.get("sources"))
    for spec in specs:
        _authorized_asset(user, spec["asset_id"])
    return {"sources": specs, "sql": _validate_sql(str(request.get("sql") or "")), "output_name": _output_filename(str(request.get("output_name") or "excel-pro-export.xlsx")), "pivot": request.get("pivot") if isinstance(request.get("pivot"), dict) else None, "chart": request.get("chart") if isinstance(request.get("chart"), dict) else None}


def _add_pivot_sheet(workbook: Workbook, connection: duckdb.DuckDBPyConnection, result_columns: list[str], config: Any):
    if not isinstance(config, dict) or not config.get("rows") or not config.get("values"):
        return None
    rows = [str(value) for value in config["rows"]]
    values = config["values"] if isinstance(config["values"], list) else []
    if any(column not in result_columns for column in rows):
        raise ExcelProError("Pivot row column does not exist in the query result")
    select = [f'"{_quote(column)}"' for column in rows]
    for item in values:
        column = str(item.get("column") or "")
        aggregate = str(item.get("aggregate") or "sum").lower()
        if column not in result_columns or aggregate not in {"sum", "count", "avg", "min", "max"}:
            raise ExcelProError("Pivot value or aggregate is invalid")
        select.append(f'{aggregate.upper()}("{_quote(column)}") AS "{_quote(aggregate + "_" + column)}"')
    query = f'SELECT {", ".join(select)} FROM excel_pro_result GROUP BY {", ".join(str(i) for i in range(1, len(rows) + 1))} ORDER BY {", ".join(str(i) for i in range(1, len(rows) + 1))}'
    cursor = connection.execute(query)
    columns = [item[0] for item in cursor.description or []]
    sheet = workbook.create_sheet("Pivot Summary")
    _write_rows(sheet, columns, cursor.fetchall())
    return sheet


def _add_chart(workbook: Workbook, result_sheet, pivot_sheet, config: Any) -> None:
    if not isinstance(config, dict):
        return
    sheet = pivot_sheet if str(config.get("sheet") or "result").lower() == "pivot" and pivot_sheet else result_sheet
    headers = [cell.value for cell in sheet[1]]
    category = str(config.get("category_column") or "")
    values = [str(value) for value in config.get("value_columns") or []]
    if category not in headers or not values or any(value not in headers for value in values) or sheet.max_row < 2:
        raise ExcelProError("Chart category/value columns are missing from the selected result sheet")
    chart_type = str(config.get("type") or "bar").lower()
    chart = {"bar": BarChart, "line": LineChart, "pie": PieChart}.get(chart_type, BarChart)()
    chart.title = str(config.get("title") or "Excel Pro Chart")[:100]
    categories = Reference(sheet, min_col=headers.index(category) + 1, min_row=2, max_row=sheet.max_row)
    for value in values:
        data = Reference(sheet, min_col=headers.index(value) + 1, min_row=1, max_row=sheet.max_row)
        chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    sheet.add_chart(chart, f"{get_column_letter(sheet.max_column + 2)}2")


def _write_rows(sheet, columns: list[str], rows: Iterable[Iterable[Any]]) -> None:
    sheet.append(columns)
    for row in rows:
        sheet.append([_xlsx_value(value) for value in row])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def _add_sources_sheet(workbook: Workbook, sql: str, citations: list[dict[str, Any]]) -> None:
    sheet = workbook.create_sheet("Sources")
    sheet.append(["DuckDB SQL", sql])
    sheet.append([])
    headers = ["Asset ID", "File", "Uploader", "Worksheet", "Range", "DuckDB Table", "Visibility", "Department Group ID"]
    sheet.append(headers)
    for item in citations:
        sheet.append([item.get("asset_id"), item.get("filename"), item.get("uploader"), item.get("sheet_name"), item.get("range"), item.get("table_name"), item.get("visibility"), item.get("department_group_id")])


def _column_names(values: list[Any]) -> list[str]:
    used = set()
    output = []
    for index, value in enumerate(values, start=1):
        base = re.sub(r"\W+", "_", str(value or "").strip(), flags=re.UNICODE).strip("_") or f"column_{index}"
        if base[0].isdigit(): base = f"_{base}"
        candidate, suffix = base, 2
        while candidate.casefold() in used:
            candidate, suffix = f"{base}_{suffix}", suffix + 1
        used.add(candidate.casefold()); output.append(candidate)
    return output


def _table_name(asset_id: int, sheet: str) -> str:
    clean = re.sub(r"\W+", "_", sheet, flags=re.UNICODE).strip("_") or "data"
    return f"asset_{asset_id}__{clean}"


def _quote(value: str) -> str:
    return value.replace('"', '""')


def _output_filename(value: str) -> str:
    name = re.sub(r"[\\/:*?\"<>|\x00-\x1f]", "_", Path(value).name).strip() or "excel-pro-export.xlsx"
    if not name.lower().endswith(".xlsx"): name += ".xlsx"
    return name[:180]


def _json_value(value: Any) -> Any:
    if isinstance(value, (date, datetime)): return value.isoformat()
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)): return None
    return value


def _xlsx_value(value: Any) -> Any:
    if isinstance(value, (str, int, float, bool, date, datetime)) or value is None: return value
    return json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else str(value)


def _json_object(value: Any) -> dict[str, Any]:
    if isinstance(value, dict): return value
    try: return json.loads(str(value or "{}"))
    except (TypeError, ValueError, json.JSONDecodeError): return {}


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""): digest.update(chunk)
    return digest.hexdigest()


def _check_cancel(job_id: int) -> None:
    if excel_pro_job_cancel_requested(job_id):
        raise InterruptedError("Excel export cancelled")


def _rpc_result(request_id: Any, result: Any) -> dict[str, Any]:
    return {"jsonrpc": "2.0", "id": request_id, "result": result}


def _rpc_error(request_id: Any, code: int, message: str) -> dict[str, Any]:
    return {"jsonrpc": "2.0", "id": request_id, "error": {"code": code, "message": message}}
