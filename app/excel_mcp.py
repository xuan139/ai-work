from __future__ import annotations

import csv
import json
import re
import sqlite3
from contextlib import contextmanager
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries

from app.db import connect, get_nas_asset


BASE_DIR = Path(__file__).resolve().parent.parent
NAS_ASSETS_DIR = BASE_DIR / "storage" / "nas_assets"
EXCEL_SUFFIXES = {".xlsx", ".xlsm", ".xltx", ".xltm"}
DELIMITED_SUFFIXES = {".csv", ".tsv"}
TABULAR_SUFFIXES = EXCEL_SUFFIXES | DELIMITED_SUFFIXES
DELIMITED_SHEET_NAME = "data"
MAX_RANGE_ROWS = 200
MAX_RANGE_COLUMNS = 50
MAX_RANGE_CELLS = 5000
MAX_SEARCH_CELLS = 100000
MAX_PROFILE_ROWS = 5000
MAX_SQL_ROWS_PER_SHEET = 20000
MAX_SQL_TOTAL_ROWS = 50000
MAX_SQL_COLUMNS = 100
MAX_SQL_RESULT_ROWS = 500

EXCEL_MCP_TOOLS = [
    {
        "name": "excel_list_workbooks",
        "title": "List NAS Spreadsheet Files",
        "description": "List Excel, CSV, and TSV files uploaded to the NAS, including asset IDs and uploader metadata.",
        "inputSchema": {
            "type": "object",
            "properties": {
                "query": {"type": "string", "description": "Optional title or filename keyword."},
                "limit": {"type": "integer", "minimum": 1, "maximum": 50, "default": 20},
            },
            "additionalProperties": False,
        },
        "annotations": {"readOnlyHint": True, "destructiveHint": False},
    },
    {
        "name": "excel_get_workbook_info",
        "title": "Get Spreadsheet Info",
        "description": "Read Excel/CSV/TSV metadata plus the SQL table and column names used by excel_query_sql.",
        "inputSchema": {
            "type": "object",
            "properties": {"asset_id": {"type": "integer", "minimum": 1}},
            "required": ["asset_id"],
            "additionalProperties": False,
        },
        "annotations": {"readOnlyHint": True, "destructiveHint": False},
    },
    {
        "name": "excel_read_range",
        "title": "Read Spreadsheet Range",
        "description": "Read values from a bounded Excel, CSV, or TSV range such as A1:F50.",
        "inputSchema": {
            "type": "object",
            "properties": {
                "asset_id": {"type": "integer", "minimum": 1},
                "sheet_name": {"type": "string"},
                "cell_range": {"type": "string", "default": "A1"},
                "max_rows": {"type": "integer", "minimum": 1, "maximum": MAX_RANGE_ROWS, "default": 50},
                "max_columns": {
                    "type": "integer",
                    "minimum": 1,
                    "maximum": MAX_RANGE_COLUMNS,
                    "default": 20,
                },
            },
            "required": ["asset_id", "sheet_name"],
            "additionalProperties": False,
        },
        "annotations": {"readOnlyHint": True, "destructiveHint": False},
    },
    {
        "name": "excel_search_cells",
        "title": "Search Spreadsheet Cells",
        "description": "Search cell values across an Excel workbook or a CSV/TSV file.",
        "inputSchema": {
            "type": "object",
            "properties": {
                "asset_id": {"type": "integer", "minimum": 1},
                "query": {"type": "string"},
                "sheet_name": {"type": "string"},
                "max_results": {"type": "integer", "minimum": 1, "maximum": 50, "default": 20},
            },
            "required": ["asset_id", "query"],
            "additionalProperties": False,
        },
        "annotations": {"readOnlyHint": True, "destructiveHint": False},
    },
    {
        "name": "excel_profile_sheet",
        "title": "Profile Spreadsheet Data",
        "description": "Inspect Excel/CSV/TSV columns, data types, missing values, distinct values, and numeric statistics.",
        "inputSchema": {
            "type": "object",
            "properties": {
                "asset_id": {"type": "integer", "minimum": 1},
                "sheet_name": {"type": "string"},
                "max_rows": {"type": "integer", "minimum": 1, "maximum": MAX_PROFILE_ROWS, "default": 1000},
            },
            "required": ["asset_id", "sheet_name"],
            "additionalProperties": False,
        },
        "annotations": {"readOnlyHint": True, "destructiveHint": False},
    },
    {
        "name": "excel_query_sql",
        "title": "Query Spreadsheets with Read-only SQL",
        "description": (
            "Load Excel worksheets or CSV/TSV data into an in-memory SQLite database and run one read-only "
            "SELECT or WITH query. Supports filtering, sorting, grouping, aggregates, and joins without "
            "changing the source file."
        ),
        "inputSchema": {
            "type": "object",
            "properties": {
                "asset_id": {"type": "integer", "minimum": 1},
                "sql": {"type": "string", "description": "One SQLite SELECT or WITH query."},
                "sheet_names": {
                    "type": "array",
                    "items": {"type": "string"},
                    "maxItems": 10,
                    "description": "Optional worksheets to load. Omit to load all worksheets.",
                },
                "max_rows": {
                    "type": "integer",
                    "minimum": 1,
                    "maximum": MAX_SQL_RESULT_ROWS,
                    "default": 100,
                },
            },
            "required": ["asset_id", "sql"],
            "additionalProperties": False,
        },
        "annotations": {"readOnlyHint": True, "destructiveHint": False},
    },
]


class ExcelMcpError(RuntimeError):
    pass


def handle_excel_mcp_request(payload: dict[str, Any]) -> dict[str, Any] | None:
    method = str(payload.get("method") or "")
    request_id = payload.get("id")
    if method == "notifications/initialized":
        return None
    if method == "initialize":
        return _result(
            request_id,
            {
                "protocolVersion": "2025-06-18",
                "capabilities": {"tools": {"listChanged": False}},
                "serverInfo": {"name": "AI Work NAS Excel/CSV Read-only MCP", "version": "1.1.0"},
                "instructions": "Read-only Excel, CSV, and TSV access for NAS-managed files. Files are never modified.",
            },
        )
    if method == "ping":
        return _result(request_id, {})
    if method == "tools/list":
        return _result(request_id, {"tools": EXCEL_MCP_TOOLS})
    if method == "tools/call":
        params = payload.get("params") if isinstance(payload.get("params"), dict) else {}
        name = str(params.get("name") or "")
        arguments = params.get("arguments") if isinstance(params.get("arguments"), dict) else {}
        try:
            output = _call_tool(name, arguments)
        except (ValueError, ExcelMcpError) as exc:
            return _error(request_id, -32602, str(exc))
        return _result(
            request_id,
            {
                "content": [{"type": "text", "text": json.dumps(output, ensure_ascii=False)}],
                "structuredContent": output,
                "isError": False,
            },
        )
    return _error(request_id, -32601, f"Method not found: {method}")


def _call_tool(name: str, arguments: dict[str, Any]) -> dict[str, Any]:
    if name == "excel_list_workbooks":
        return _list_workbooks(str(arguments.get("query") or "").strip(), _bounded_int(arguments.get("limit"), 20, 50))
    if name == "excel_get_workbook_info":
        asset, path = _workbook_asset(arguments)
        if _is_delimited(path):
            sheets = [_delimited_info(path)]
        else:
            workbook = _open_workbook(path)
            try:
                sheets = []
                used_table_names: set[str] = set()
                for index, sheet in enumerate(workbook.worksheets, start=1):
                    first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
                    sheets.append(
                        {
                            "name": sheet.title,
                            "state": sheet.sheet_state,
                            "max_row": sheet.max_row,
                            "max_column": sheet.max_column,
                            "dimension": sheet.calculate_dimension(),
                            "sql_table_name": _unique_identifier(sheet.title, used_table_names, f"sheet_{index}"),
                            "sql_columns": _column_names(first_row[:MAX_SQL_COLUMNS]),
                        }
                    )
            finally:
                workbook.close()
        return {"workbook": _asset_summary(asset), "sheet_count": len(sheets), "sheets": sheets}
    if name == "excel_read_range":
        asset, path = _workbook_asset(arguments)
        if _is_delimited(path):
            return {"workbook": _asset_summary(asset), **_read_delimited_range(path, arguments)}
        workbook = _open_workbook(path)
        try:
            sheet = _worksheet(workbook, arguments)
            min_column, min_row, max_column, max_row = _range_bounds(arguments, sheet)
            rows = [
                [_json_value(cell.value) for cell in row]
                for row in sheet.iter_rows(
                    min_row=min_row,
                    max_row=max_row,
                    min_col=min_column,
                    max_col=max_column,
                )
            ]
            sheet_name = sheet.title
            range_label = (
                f"{get_column_letter(min_column)}{min_row}:"
                f"{get_column_letter(max_column)}{max_row}"
            )
        finally:
            workbook.close()
        return {
            "workbook": _asset_summary(asset),
            "sheet_name": sheet_name,
            "range": range_label,
            "row_count": len(rows),
            "column_count": max_column - min_column + 1,
            "rows": rows,
        }
    if name == "excel_search_cells":
        asset, path = _workbook_asset(arguments)
        query = str(arguments.get("query") or "").strip()
        if not query:
            raise ValueError("query is required")
        limit = _bounded_int(arguments.get("max_results"), 20, 50)
        if _is_delimited(path):
            return {
                "workbook": _asset_summary(asset),
                "query": query,
                **_search_delimited(path, arguments, query, limit),
            }
        workbook = _open_workbook(path)
        try:
            requested_sheet = str(arguments.get("sheet_name") or "").strip()
            sheets = [_worksheet(workbook, {"sheet_name": requested_sheet})] if requested_sheet else workbook.worksheets
            matches, scanned, truncated = _search_sheets(sheets, query, limit)
        finally:
            workbook.close()
        return {
            "workbook": _asset_summary(asset),
            "query": query,
            "count": len(matches),
            "scanned_cells": scanned,
            "scan_truncated": truncated,
            "matches": matches,
        }
    if name == "excel_profile_sheet":
        asset, path = _workbook_asset(arguments)
        if _is_delimited(path):
            return {
                "workbook": _asset_summary(asset),
                **_profile_delimited(
                    path,
                    arguments,
                    _bounded_int(arguments.get("max_rows"), 1000, MAX_PROFILE_ROWS),
                ),
            }
        workbook = _open_workbook(path)
        try:
            sheet = _worksheet(workbook, arguments)
            profile = _profile_sheet(
                sheet,
                _bounded_int(arguments.get("max_rows"), 1000, MAX_PROFILE_ROWS),
            )
        finally:
            workbook.close()
        return {"workbook": _asset_summary(asset), **profile}
    if name == "excel_query_sql":
        asset, path = _workbook_asset(arguments)
        sql = str(arguments.get("sql") or "").strip()
        if not sql:
            raise ValueError("sql is required")
        sheet_names = arguments.get("sheet_names")
        if sheet_names is not None and not isinstance(sheet_names, list):
            raise ValueError("sheet_names must be an array")
        result = _query_sql(
            path,
            sql,
            [str(item) for item in sheet_names] if sheet_names else None,
            _bounded_int(arguments.get("max_rows"), 100, MAX_SQL_RESULT_ROWS),
        )
        return {"workbook": _asset_summary(asset), **result}
    raise ValueError(f"Unknown tool: {name}")


def _list_workbooks(query: str, limit: int) -> dict[str, Any]:
    suffix_sql = " OR ".join("lower(nas_assets.original_filename) LIKE ?" for _ in TABULAR_SUFFIXES)
    params: list[Any] = [f"%{suffix}" for suffix in sorted(TABULAR_SUFFIXES)]
    query_sql = ""
    if query:
        query_sql = "AND (nas_assets.title LIKE ? OR nas_assets.original_filename LIKE ?)"
        params.extend([f"%{query}%", f"%{query}%"])
    params.append(limit)
    with connect() as conn:
        rows = conn.execute(
            f"""
            SELECT nas_assets.*, users.username AS owner_username
            FROM nas_assets
            JOIN users ON users.id = nas_assets.user_id
            WHERE ({suffix_sql}) {query_sql}
            ORDER BY datetime(nas_assets.created_at) DESC, nas_assets.id DESC
            LIMIT ?
            """,
            params,
        ).fetchall()
    workbooks = [_asset_summary(dict(row)) for row in rows]
    return {"query": query or None, "count": len(workbooks), "workbooks": workbooks}


def _workbook_asset(arguments: dict[str, Any]) -> tuple[dict[str, Any], Path]:
    try:
        asset_id = int(arguments.get("asset_id"))
    except (TypeError, ValueError) as exc:
        raise ValueError("asset_id is required") from exc
    asset = get_nas_asset(asset_id)
    if not asset:
        raise ValueError("Spreadsheet asset not found")
    path = Path(str(asset.get("stored_path") or "")).resolve()
    if path.suffix.lower() not in TABULAR_SUFFIXES:
        raise ValueError("The selected NAS asset is not an .xlsx, .xlsm, .csv, or .tsv file")
    try:
        path.relative_to(NAS_ASSETS_DIR.resolve())
    except ValueError as exc:
        raise ExcelMcpError("Spreadsheet is outside the managed NAS asset directory") from exc
    if not path.is_file():
        raise ExcelMcpError("Spreadsheet file is missing from NAS storage")
    return asset, path


def _open_workbook(path: Path):
    try:
        return load_workbook(path, read_only=True, data_only=True, keep_links=False)
    except Exception as exc:
        raise ExcelMcpError(f"Unable to read Excel workbook: {exc}") from exc


def _is_delimited(path: Path) -> bool:
    return path.suffix.lower() in DELIMITED_SUFFIXES


@contextmanager
def _open_delimited(path: Path):
    encoding = _detect_delimited_encoding(path)
    handle = path.open("r", encoding=encoding, newline="")
    try:
        sample = handle.read(65536)
        handle.seek(0)
        delimiter = _detect_delimiter(sample, path.suffix.lower())
        yield csv.reader(handle, delimiter=delimiter), {
            "encoding": encoding,
            "delimiter": delimiter,
            "delimiter_name": "tab" if delimiter == "\t" else delimiter,
        }
    finally:
        handle.close()


def _detect_delimited_encoding(path: Path) -> str:
    with path.open("rb") as handle:
        sample = handle.read(65536)
    if sample.startswith((b"\xff\xfe", b"\xfe\xff")):
        return "utf-16"
    for encoding in ("utf-8-sig", "cp950"):
        try:
            sample.decode(encoding)
            return encoding
        except UnicodeDecodeError:
            continue
    raise ExcelMcpError("CSV/TSV encoding is not supported; use UTF-8, UTF-16, Big5, or CP950")


def _detect_delimiter(sample: str, suffix: str) -> str:
    fallback = "\t" if suffix == ".tsv" else ","
    if not sample.strip():
        return fallback
    try:
        return csv.Sniffer().sniff(sample, delimiters=",\t;|").delimiter
    except csv.Error:
        return fallback


def _require_delimited_sheet(arguments: dict[str, Any], path: Path) -> None:
    sheet_name = str(arguments.get("sheet_name") or "").strip()
    allowed = {DELIMITED_SHEET_NAME, path.stem}
    if sheet_name and sheet_name not in allowed:
        raise ValueError(f"CSV/TSV has one data sheet named '{DELIMITED_SHEET_NAME}'")


def _delimited_info(path: Path) -> dict[str, Any]:
    with _open_delimited(path) as (reader, metadata):
        first_row = next(reader, [])
        max_column = len(first_row)
        data_rows = 0
        truncated = False
        for row in reader:
            if data_rows >= MAX_SQL_TOTAL_ROWS:
                truncated = True
                break
            data_rows += 1
            max_column = max(max_column, len(row))
    max_row = data_rows + (1 if first_row else 0)
    dimension = "A1"
    if max_row and max_column:
        dimension = f"A1:{get_column_letter(max_column)}{max_row}"
    return {
        "name": DELIMITED_SHEET_NAME,
        "state": "visible",
        "max_row": max_row,
        "max_column": max_column,
        "dimension": dimension,
        "row_count_truncated": truncated,
        "sql_table_name": _unique_identifier(path.stem, set(), "data"),
        "sql_columns": _column_names(first_row[:MAX_SQL_COLUMNS]),
        **metadata,
    }


def _read_delimited_range(path: Path, arguments: dict[str, Any]) -> dict[str, Any]:
    _require_delimited_sheet(arguments, path)
    min_column, min_row, max_column, max_row = _delimited_range_bounds(arguments)
    rows: list[list[Any]] = []
    with _open_delimited(path) as (reader, metadata):
        for row_number, row in enumerate(reader, start=1):
            if row_number < min_row:
                continue
            if row_number > max_row:
                break
            values = row[min_column - 1:max_column]
            values.extend([""] * (max_column - min_column + 1 - len(values)))
            rows.append([_parse_delimited_value(value) for value in values])
    return {
        "sheet_name": DELIMITED_SHEET_NAME,
        "range": f"{get_column_letter(min_column)}{min_row}:{get_column_letter(max_column)}{max_row}",
        "row_count": len(rows),
        "column_count": max_column - min_column + 1,
        "rows": rows,
        **metadata,
    }


def _delimited_range_bounds(arguments: dict[str, Any]) -> tuple[int, int, int, int]:
    cell_range = str(arguments.get("cell_range") or "A1").strip().upper()
    if "!" in cell_range:
        raise ValueError("Use sheet_name separately from cell_range")
    try:
        min_column, min_row, requested_max_column, requested_max_row = range_boundaries(cell_range)
    except ValueError as exc:
        raise ValueError("cell_range must use A1 notation, for example A1:F50") from exc
    max_rows = _bounded_int(arguments.get("max_rows"), 50, MAX_RANGE_ROWS)
    max_columns = _bounded_int(arguments.get("max_columns"), 20, MAX_RANGE_COLUMNS)
    if ":" not in cell_range:
        requested_max_row = min_row + max_rows - 1
        requested_max_column = min_column + max_columns - 1
    max_row = min(requested_max_row, min_row + max_rows - 1)
    max_column = min(requested_max_column, min_column + max_columns - 1)
    if (max_row - min_row + 1) * (max_column - min_column + 1) > MAX_RANGE_CELLS:
        raise ValueError(f"Requested range exceeds {MAX_RANGE_CELLS} cells")
    return min_column, min_row, max_column, max_row


def _search_delimited(
    path: Path,
    arguments: dict[str, Any],
    query: str,
    limit: int,
) -> dict[str, Any]:
    _require_delimited_sheet(arguments, path)
    needle = query.casefold()
    matches: list[dict[str, Any]] = []
    scanned = 0
    truncated = False
    with _open_delimited(path) as (reader, metadata):
        for row_number, row in enumerate(reader, start=1):
            for column_number, value in enumerate(row, start=1):
                scanned += 1
                if needle in value.casefold():
                    matches.append(
                        {
                            "sheet_name": DELIMITED_SHEET_NAME,
                            "cell": f"{get_column_letter(column_number)}{row_number}",
                            "value": _parse_delimited_value(value),
                        }
                    )
                    if len(matches) >= limit:
                        return {
                            "count": len(matches),
                            "scanned_cells": scanned,
                            "scan_truncated": True,
                            "matches": matches,
                            **metadata,
                        }
                if scanned >= MAX_SEARCH_CELLS:
                    truncated = True
                    break
            if truncated:
                break
    return {
        "count": len(matches),
        "scanned_cells": scanned,
        "scan_truncated": truncated,
        "matches": matches,
        **metadata,
    }


def _worksheet(workbook, arguments: dict[str, Any]):
    sheet_name = str(arguments.get("sheet_name") or "").strip()
    if not sheet_name:
        raise ValueError("sheet_name is required")
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Worksheet not found: {sheet_name}")
    return workbook[sheet_name]


def _range_bounds(arguments: dict[str, Any], sheet) -> tuple[int, int, int, int]:
    cell_range = str(arguments.get("cell_range") or "A1").strip().upper()
    if "!" in cell_range:
        raise ValueError("Use sheet_name separately from cell_range")
    try:
        min_column, min_row, requested_max_column, requested_max_row = range_boundaries(cell_range)
    except ValueError as exc:
        raise ValueError("cell_range must use A1 notation, for example A1:F50") from exc
    max_rows = _bounded_int(arguments.get("max_rows"), 50, MAX_RANGE_ROWS)
    max_columns = _bounded_int(arguments.get("max_columns"), 20, MAX_RANGE_COLUMNS)
    if ":" not in cell_range:
        requested_max_row = min(sheet.max_row, min_row + max_rows - 1)
        requested_max_column = min(sheet.max_column, min_column + max_columns - 1)
    max_row = min(requested_max_row, min_row + max_rows - 1)
    max_column = min(requested_max_column, min_column + max_columns - 1)
    if (max_row - min_row + 1) * (max_column - min_column + 1) > MAX_RANGE_CELLS:
        raise ValueError(f"Requested range exceeds {MAX_RANGE_CELLS} cells")
    return min_column, min_row, max_column, max_row


def _search_sheets(sheets: list[Any], query: str, limit: int) -> tuple[list[dict[str, Any]], int, bool]:
    needle = query.casefold()
    matches: list[dict[str, Any]] = []
    scanned = 0
    for sheet in sheets:
        for row in sheet.iter_rows():
            for cell in row:
                scanned += 1
                if cell.value is not None and needle in str(cell.value).casefold():
                    matches.append(
                        {"sheet_name": sheet.title, "cell": cell.coordinate, "value": _json_value(cell.value)}
                    )
                    if len(matches) >= limit:
                        return matches, scanned, scanned >= MAX_SEARCH_CELLS
                if scanned >= MAX_SEARCH_CELLS:
                    return matches, scanned, True
    return matches, scanned, False


def _profile_sheet(sheet: Any, max_rows: int) -> dict[str, Any]:
    iterator = sheet.iter_rows(values_only=True)
    first_row = next(iterator, ())
    return _profile_values(
        sheet.title,
        first_row,
        iterator,
        max_rows,
        max(0, sheet.max_row - 1),
    )


def _profile_delimited(path: Path, arguments: dict[str, Any], max_rows: int) -> dict[str, Any]:
    _require_delimited_sheet(arguments, path)
    with _open_delimited(path) as (reader, metadata):
        first_row = next(reader, [])
        values = ([_parse_delimited_value(value) for value in row] for row in reader)
        result = _profile_values(DELIMITED_SHEET_NAME, first_row, values, max_rows, None)
    return {**result, **metadata}


def _profile_values(
    sheet_name: str,
    first_row: tuple[Any, ...] | list[Any],
    rows: Any,
    max_rows: int,
    worksheet_rows: int | None,
) -> dict[str, Any]:
    column_count = min(len(first_row), MAX_SQL_COLUMNS)
    columns = _column_names(first_row[:column_count])
    stats = [
        {
            "column": column,
            "non_empty_count": 0,
            "empty_count": 0,
            "distinct_values": set(),
            "types": {},
            "numeric_count": 0,
            "numeric_sum": 0.0,
            "numeric_min": None,
            "numeric_max": None,
        }
        for column in columns
    ]
    scanned_rows = 0
    scan_truncated = False
    for row in rows:
        if scanned_rows >= max_rows:
            scan_truncated = True
            break
        scanned_rows += 1
        for index, column_stats in enumerate(stats):
            value = row[index] if index < len(row) else None
            if value is None or value == "":
                column_stats["empty_count"] += 1
                continue
            column_stats["non_empty_count"] += 1
            value_type = _value_type(value)
            column_stats["types"][value_type] = column_stats["types"].get(value_type, 0) + 1
            if len(column_stats["distinct_values"]) < 1000:
                column_stats["distinct_values"].add(str(_json_value(value)))
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                numeric = float(value)
                column_stats["numeric_count"] += 1
                column_stats["numeric_sum"] += numeric
                current_min = column_stats["numeric_min"]
                current_max = column_stats["numeric_max"]
                column_stats["numeric_min"] = numeric if current_min is None else min(current_min, numeric)
                column_stats["numeric_max"] = numeric if current_max is None else max(current_max, numeric)

    output_columns = []
    for column_stats in stats:
        numeric_count = column_stats.pop("numeric_count")
        numeric_sum = column_stats.pop("numeric_sum")
        distinct_values = column_stats.pop("distinct_values")
        output_columns.append(
            {
                **column_stats,
                "distinct_count": len(distinct_values),
                "distinct_count_capped": len(distinct_values) >= 1000,
                "numeric_average": round(numeric_sum / numeric_count, 6) if numeric_count else None,
            }
        )
    return {
        "sheet_name": sheet_name,
        "worksheet_rows": worksheet_rows,
        "scanned_rows": scanned_rows,
        "scan_truncated": scan_truncated if worksheet_rows is None else worksheet_rows > scanned_rows,
        "columns": output_columns,
    }


def _query_sql(path: Path, sql: str, sheet_names: list[str] | None, max_rows: int) -> dict[str, Any]:
    if not re.match(r"^(SELECT|WITH)\b", sql, flags=re.IGNORECASE):
        raise ValueError("Only one read-only SELECT or WITH query is allowed")

    workbook = None
    connection = sqlite3.connect(":memory:")
    connection.row_factory = sqlite3.Row
    try:
        if _is_delimited(path):
            _validate_delimited_sheet_names(path, sheet_names)
            table_map, loaded_rows, truncated = _load_delimited_sql_table(connection, path)
        else:
            workbook = _open_workbook(path)
            selected_sheets = _selected_sql_sheets(workbook, sheet_names)
            table_map, loaded_rows, truncated = _load_sql_tables(connection, selected_sheets)
        connection.set_authorizer(_readonly_sql_authorizer)
        try:
            cursor = connection.execute(sql)
            rows = cursor.fetchmany(max_rows + 1)
        except sqlite3.Error as exc:
            available = ", ".join(item["table_name"] for item in table_map)
            raise ValueError(f"SQL query failed: {exc}. Available tables: {available}") from exc
        result_truncated = len(rows) > max_rows
        rows = rows[:max_rows]
        columns = [item[0] for item in cursor.description or []]
        return {
            "sql": sql,
            "tables": table_map,
            "loaded_rows": loaded_rows,
            "source_truncated": truncated,
            "columns": columns,
            "row_count": len(rows),
            "result_truncated": result_truncated,
            "rows": [[_json_value(value) for value in row] for row in rows],
        }
    finally:
        connection.close()
        if workbook is not None:
            workbook.close()


def _validate_delimited_sheet_names(path: Path, sheet_names: list[str] | None) -> None:
    if not sheet_names:
        return
    allowed = {DELIMITED_SHEET_NAME, path.stem}
    missing = [name for name in sheet_names if name not in allowed]
    if missing:
        raise ValueError(f"CSV/TSV has one data sheet named '{DELIMITED_SHEET_NAME}'")


def _load_delimited_sql_table(
    connection: sqlite3.Connection,
    path: Path,
) -> tuple[list[dict[str, Any]], int, bool]:
    table_name = _unique_identifier(path.stem, set(), "data")
    with _open_delimited(path) as (reader, metadata):
        first_row = next(reader, [])
        column_count = min(len(first_row), MAX_SQL_COLUMNS)
        if not column_count:
            return ([{
                "sheet_name": DELIMITED_SHEET_NAME,
                "table_name": table_name,
                "columns": [],
                "rows": 0,
                **metadata,
            }], 0, False)
        columns = _column_names(first_row[:column_count])
        quoted_columns = ", ".join(f'"{_quote_identifier(column)}"' for column in columns)
        connection.execute(f'CREATE TABLE "{_quote_identifier(table_name)}" ({quoted_columns})')
        placeholders = ", ".join("?" for _ in columns)
        insert_sql = f'INSERT INTO "{_quote_identifier(table_name)}" VALUES ({placeholders})'
        loaded_rows = 0
        truncated = False
        for row in reader:
            if loaded_rows >= MAX_SQL_TOTAL_ROWS:
                truncated = True
                break
            values = [
                _parse_delimited_value(row[column]) if column < len(row) else None
                for column in range(column_count)
            ]
            connection.execute(insert_sql, values)
            loaded_rows += 1
    return ([{
        "sheet_name": DELIMITED_SHEET_NAME,
        "table_name": table_name,
        "columns": columns,
        "rows": loaded_rows,
        **metadata,
    }], loaded_rows, truncated)


def _selected_sql_sheets(workbook: Any, sheet_names: list[str] | None) -> list[Any]:
    if not sheet_names:
        return workbook.worksheets
    missing = [name for name in sheet_names if name not in workbook.sheetnames]
    if missing:
        raise ValueError(f"Worksheet not found: {', '.join(missing)}")
    return [workbook[name] for name in dict.fromkeys(sheet_names)]


def _load_sql_tables(connection: sqlite3.Connection, sheets: list[Any]) -> tuple[list[dict[str, Any]], int, bool]:
    used_table_names: set[str] = set()
    table_map: list[dict[str, Any]] = []
    total_rows = 0
    source_truncated = False
    for index, sheet in enumerate(sheets, start=1):
        table_name = _unique_identifier(sheet.title, used_table_names, f"sheet_{index}")
        iterator = sheet.iter_rows(values_only=True)
        first_row = next(iterator, ())
        column_count = min(len(first_row), MAX_SQL_COLUMNS)
        if not column_count:
            table_map.append({"sheet_name": sheet.title, "table_name": table_name, "columns": [], "rows": 0})
            continue
        columns = _column_names(first_row[:column_count])
        quoted_columns = ", ".join(f'"{_quote_identifier(column)}"' for column in columns)
        connection.execute(f'CREATE TABLE "{_quote_identifier(table_name)}" ({quoted_columns})')
        placeholders = ", ".join("?" for _ in columns)
        insert_sql = f'INSERT INTO "{_quote_identifier(table_name)}" VALUES ({placeholders})'
        sheet_rows = 0
        for row in iterator:
            if sheet_rows >= MAX_SQL_ROWS_PER_SHEET or total_rows >= MAX_SQL_TOTAL_ROWS:
                source_truncated = True
                break
            values = [_sqlite_value(row[column]) if column < len(row) else None for column in range(column_count)]
            connection.execute(insert_sql, values)
            sheet_rows += 1
            total_rows += 1
        table_map.append(
            {
                "sheet_name": sheet.title,
                "table_name": table_name,
                "columns": columns,
                "rows": sheet_rows,
            }
        )
        if total_rows >= MAX_SQL_TOTAL_ROWS:
            source_truncated = True
            break
    return table_map, total_rows, source_truncated


def _column_names(values: tuple[Any, ...] | list[Any]) -> list[str]:
    used: set[str] = set()
    return [
        _unique_identifier(str(value).strip() if value not in (None, "") else "", used, f"column_{index}")
        for index, value in enumerate(values, start=1)
    ]


def _unique_identifier(value: str, used: set[str], fallback: str) -> str:
    base = re.sub(r"\W+", "_", value, flags=re.UNICODE).strip("_") or fallback
    if base[0].isdigit():
        base = f"_{base}"
    candidate = base
    suffix = 2
    while candidate.casefold() in used:
        candidate = f"{base}_{suffix}"
        suffix += 1
    used.add(candidate.casefold())
    return candidate


def _quote_identifier(value: str) -> str:
    return value.replace('"', '""')


def _parse_delimited_value(value: str) -> Any:
    normalized = value.strip()
    if not normalized:
        return None
    if re.fullmatch(r"[-+]?(?:0|[1-9]\d*)", normalized):
        try:
            return int(normalized)
        except ValueError:
            pass
    if re.fullmatch(r"[-+]?(?:\d+\.\d*|\.\d+)(?:[eE][-+]?\d+)?", normalized):
        try:
            return float(normalized)
        except ValueError:
            pass
    return normalized


def _sqlite_value(value: Any) -> Any:
    converted = _json_value(value)
    if isinstance(converted, (str, int, float, bytes)) or converted is None:
        return converted
    if isinstance(converted, bool):
        return int(converted)
    return str(converted)


def _value_type(value: Any) -> str:
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, (int, float, Decimal)):
        return "number"
    if isinstance(value, (datetime, date, time)):
        return "date_time"
    return "text"


def _readonly_sql_authorizer(action: int, _arg1: str, _arg2: str, _database: str, _trigger: str) -> int:
    denied_actions = {
        sqlite3.SQLITE_INSERT,
        sqlite3.SQLITE_UPDATE,
        sqlite3.SQLITE_DELETE,
        sqlite3.SQLITE_CREATE_INDEX,
        sqlite3.SQLITE_CREATE_TABLE,
        sqlite3.SQLITE_CREATE_TEMP_INDEX,
        sqlite3.SQLITE_CREATE_TEMP_TABLE,
        sqlite3.SQLITE_CREATE_TEMP_TRIGGER,
        sqlite3.SQLITE_CREATE_TEMP_VIEW,
        sqlite3.SQLITE_CREATE_TRIGGER,
        sqlite3.SQLITE_CREATE_VIEW,
        sqlite3.SQLITE_DROP_INDEX,
        sqlite3.SQLITE_DROP_TABLE,
        sqlite3.SQLITE_DROP_TEMP_INDEX,
        sqlite3.SQLITE_DROP_TEMP_TABLE,
        sqlite3.SQLITE_DROP_TEMP_TRIGGER,
        sqlite3.SQLITE_DROP_TEMP_VIEW,
        sqlite3.SQLITE_DROP_TRIGGER,
        sqlite3.SQLITE_DROP_VIEW,
        sqlite3.SQLITE_ALTER_TABLE,
        sqlite3.SQLITE_REINDEX,
        sqlite3.SQLITE_ANALYZE,
        sqlite3.SQLITE_ATTACH,
        sqlite3.SQLITE_DETACH,
        sqlite3.SQLITE_PRAGMA,
        sqlite3.SQLITE_TRANSACTION,
        sqlite3.SQLITE_SAVEPOINT,
    }
    return sqlite3.SQLITE_DENY if action in denied_actions else sqlite3.SQLITE_OK


def _asset_summary(asset: dict[str, Any]) -> dict[str, Any]:
    return {
        "asset_id": asset.get("id"),
        "title": asset.get("title"),
        "filename": asset.get("original_filename"),
        "file_size": asset.get("file_size"),
        "uploader": asset.get("owner_username"),
        "uploaded_at": asset.get("created_at"),
    }


def _json_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    return str(value)


def _bounded_int(value: object, default: int, maximum: int) -> int:
    try:
        number = int(value) if value is not None else default
    except (TypeError, ValueError) as exc:
        raise ValueError("Numeric limits must be integers") from exc
    return max(1, min(number, maximum))


def _result(request_id: object, result: dict[str, Any]) -> dict[str, Any]:
    return {"jsonrpc": "2.0", "id": request_id, "result": result}


def _error(request_id: object, code: int, message: str) -> dict[str, Any]:
    return {"jsonrpc": "2.0", "id": request_id, "error": {"code": code, "message": message}}
