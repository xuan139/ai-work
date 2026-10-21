import hashlib
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from app import db, excel_pro
from app.auth import create_mcp_delegation_headers, hash_password, read_mcp_delegated_user


class ExcelProTests(unittest.TestCase):
    def setUp(self) -> None:
        self.directory = tempfile.TemporaryDirectory()
        self.root = Path(self.directory.name)
        self.assets_dir = self.root / "nas_assets"
        self.assets_dir.mkdir()
        self.db_patch = patch.object(db, "DB_PATH", self.root / "app.db")
        self.output_patch = patch.object(excel_pro, "NAS_ASSETS_DIR", self.assets_dir)
        self.db_patch.start()
        self.output_patch.start()
        db.init_db()
        db.seed_admin(hash_password("admin123"))
        self.admin = db.get_user_by_username("admin")
        self.sales_user = db.create_user(username="sales", password_hash=hash_password("test123"), role="user")
        self.outsider = db.create_user(username="outside", password_hash=hash_password("test123"), role="user")
        self.group = db.create_user_group(name="Sales Department", created_by=self.admin["id"])
        db.set_user_group_members(self.group["id"], [self.sales_user["id"]])

        self.sales_path = self.assets_dir / "sales.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Sales Detail"
        sheet.append(["Product", "Quantity", "Revenue"])
        sheet.append(["NAS A", 3, 300])
        sheet.append(["NAS B", 2, 300])
        sheet.append(["NAS A", 1, 100])
        workbook.save(self.sales_path)
        workbook.close()
        self.sales_asset = db.create_nas_asset(
            user_id=self.sales_user["id"], category="excel", title="Sales",
            original_filename="sales.xlsx", stored_path=str(self.sales_path),
            mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            file_size=self.sales_path.stat().st_size, status="completed", visibility="private",
        )

        self.inventory_path = self.assets_dir / "inventory.csv"
        self.inventory_path.write_text("Product,Available\nNAS A,10\nNAS B,1\n", encoding="utf-8")
        self.inventory_asset = db.create_nas_asset(
            user_id=self.admin["id"], category="excel", title="Inventory",
            original_filename="inventory.csv", stored_path=str(self.inventory_path),
            mime_type="text/csv", file_size=self.inventory_path.stat().st_size,
            status="completed", visibility="group", owner_group_id=self.group["id"],
        )

    def tearDown(self) -> None:
        self.output_patch.stop()
        self.db_patch.stop()
        self.directory.cleanup()

    def source_specs(self):
        return [{"asset_id": self.sales_asset["id"]}, {"asset_id": self.inventory_asset["id"]}]

    def query(self):
        sales = f'asset_{self.sales_asset["id"]}__Sales_Detail'
        inventory = f'asset_{self.inventory_asset["id"]}__data'
        return (
            f'SELECT s.Product, SUM(s.Quantity) AS Total_Quantity, MAX(i.Available) AS Available '
            f'FROM "{sales}" s JOIN "{inventory}" i ON i.Product = s.Product '
            f'GROUP BY s.Product ORDER BY s.Product'
        )

    def digest(self, path):
        return hashlib.sha256(path.read_bytes()).hexdigest()

    def test_lists_only_sources_allowed_by_owner_or_department(self) -> None:
        allowed = excel_pro.list_sources(self.sales_user)
        self.assertEqual({row["asset_id"] for row in allowed["sources"]}, {self.sales_asset["id"], self.inventory_asset["id"]})
        denied = excel_pro.list_sources(self.outsider)
        self.assertEqual(denied["sources"], [])
        with self.assertRaises(excel_pro.ExcelProError):
            excel_pro.describe_sources(self.outsider, [{"asset_id": self.sales_asset["id"]}])

    def test_cross_file_duckdb_query_returns_precise_citations(self) -> None:
        result = excel_pro.run_query(self.sales_user, self.source_specs(), self.query(), 100)
        self.assertEqual(result["columns"], ["Product", "Total_Quantity", "Available"])
        self.assertEqual(result["rows"], [["NAS A", 4, 10], ["NAS B", 2, 1]])
        self.assertEqual({item["range"] for item in result["citations"]}, {"Sales Detail!A2:C4", "data!A2:B3"})
        self.assertTrue(result["engine"].startswith("DuckDB"))

    def test_export_creates_new_xlsx_with_pivot_chart_and_sources(self) -> None:
        before = {self.sales_path: self.digest(self.sales_path), self.inventory_path: self.digest(self.inventory_path)}
        request = {
            "sources": self.source_specs(), "sql": self.query(), "output_name": "analysis.xlsx",
            "pivot": {"rows": ["Product"], "values": [{"column": "Total_Quantity", "aggregate": "sum"}]},
            "chart": {"type": "bar", "sheet": "pivot", "category_column": "Product", "value_columns": ["sum_Total_Quantity"], "title": "Quantity"},
        }
        queued = excel_pro.enqueue_export(self.sales_user, request)
        job = db.get_excel_pro_job(queued["id"], user_id=self.sales_user["id"], role="user")
        result = excel_pro.execute_export_job(job, lambda value, stage: None)
        output = db.get_nas_asset(result["output_asset_id"])
        self.assertEqual(output["user_id"], self.sales_user["id"])
        self.assertEqual(output["visibility"], "private")
        self.assertNotIn(Path(output["stored_path"]), before)
        workbook = load_workbook(output["stored_path"])
        self.assertEqual(workbook.sheetnames, ["Query Result", "Pivot Summary", "Sources"])
        self.assertEqual(len(workbook["Pivot Summary"]._charts), 1)
        self.assertEqual(workbook["Sources"]["E4"].value, "Sales Detail!A2:C4")
        workbook.close()
        self.assertEqual(before, {path: self.digest(path) for path in before})

    def test_mcp_requires_user_delegation_for_calls(self) -> None:
        request = {"jsonrpc": "2.0", "id": 1, "method": "tools/call", "params": {"name": "excel_pro_list_sources", "arguments": {}}}
        denied = excel_pro.handle_excel_pro_mcp_request(request, None)
        self.assertEqual(denied["error"]["code"], -32001)
        allowed = excel_pro.handle_excel_pro_mcp_request(request, self.sales_user)
        self.assertEqual(allowed["result"]["structuredContent"]["count"], 2)

    def test_sql_file_access_and_mutation_are_blocked(self) -> None:
        for sql in ("DELETE FROM x", "SELECT * FROM read_csv('/etc/passwd')", "SELECT 1; SELECT 2"):
            with self.assertRaises(excel_pro.ExcelProError):
                excel_pro.run_query(self.sales_user, self.source_specs(), sql)

    def test_tool_annotations_separate_query_from_export(self) -> None:
        tools = {tool["name"]: tool for tool in excel_pro.EXCEL_PRO_TOOLS}
        self.assertTrue(tools["excel_pro_query"]["annotations"]["readOnlyHint"])
        self.assertFalse(tools["excel_pro_start_export"]["annotations"]["readOnlyHint"])

    def test_signed_mcp_delegation_resolves_active_database_user(self) -> None:
        headers = create_mcp_delegation_headers(self.sales_user)
        delegated = read_mcp_delegated_user(headers)
        self.assertEqual(delegated["id"], self.sales_user["id"])
        headers["X-AI-Work-User-Id"] = str(self.outsider["id"])
        self.assertIsNone(read_mcp_delegated_user(headers))


if __name__ == "__main__":
    unittest.main()
